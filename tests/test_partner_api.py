"""API đối tác /v1 (M4): xác thực, cách ly tenant, sandbox 0 lần gọi LLM, idempotency, chống trùng,
hạn mức, webhook ký HMAC + retry, retention, không ghi nội dung PDF vào log.

SQLite + provider mock, worker chạy inline, webhook gửi qua httpx.MockTransport. Không gọi API thật.
"""

import io
import json
import logging
import uuid
import zipfile
from collections.abc import AsyncIterator, Iterator
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.api.main import create_app
from app.api.routes.internal import get_session
from app.core.config import Settings, get_settings
from app.db import models
from app.db.session import Base
from app.providers.base import ExtractionProvider, ExtractionRequest, ExtractionResult
from app.providers.mock import MockProvider
from app.services import retention, webhooks
from app.services.api_keys import new_api_key
from app.services.documents import process_document
from app.services.ratelimit import MemoryRateLimiter
from app.web.auth import DEV_COOKIE
from tests.helpers import images_to_pdf, sample_extraction, text_page

ADMIN = "caonntt@gmail.com"
TENANT_A = uuid.UUID("00000000-0000-0000-0000-0000000000aa")  # require_human_review=true
TENANT_B = uuid.UUID("00000000-0000-0000-0000-0000000000bb")  # require_human_review=false
SECRET_A = "whsec_test_a"


class GuardProvider(ExtractionProvider):
    """Bọc provider mock, đếm số lần gọi (sandbox phải = 0)."""

    name = "mock"

    def __init__(self, inner: MockProvider) -> None:
        self.inner = inner
        self.model = "mock"

    @property
    def calls(self) -> list[ExtractionRequest]:
        return self.inner.calls

    async def extract(self, req: ExtractionRequest) -> ExtractionResult:
        return await self.inner.extract(req)

    async def health(self) -> bool:
        return True


class Env(SimpleNamespace):
    client: TestClient
    settings: Settings
    provider: GuardProvider
    maker: async_sessionmaker[AsyncSession]
    keys: dict[str, str]
    hooks: list[httpx.Request]
    hook_status: list[int]

    def h(self, key: str, **extra: str) -> dict[str, str]:
        return {"Authorization": f"Bearer {self.keys[key]}", **extra}

    def db(self, fn):  # type: ignore[no-untyped-def]
        async def run():  # type: ignore[no-untyped-def]
            async with self.maker() as s:
                return await fn(s)

        return self.client.portal.call(run)  # type: ignore[union-attr]

    def deliver(self) -> int:
        async def run(s: AsyncSession) -> int:
            def handler(req: httpx.Request) -> httpx.Response:
                self.hooks.append(req)
                return httpx.Response(self.hook_status.pop(0) if self.hook_status else 200)

            async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as c:
                return await webhooks.deliver_due(s, self.settings, c)

        return self.db(run)  # type: ignore[no-any-return]

    def upload(self, key: str, files: list[tuple[str, bytes]], **data: Any) -> httpx.Response:
        headers = self.h(key)
        if "idem" in data:
            headers["Idempotency-Key"] = data.pop("idem")
        return self.client.post(
            "/v1/documents",
            headers=headers,
            files=[("files", (n, b, "application/pdf")) for n, b in files],
            data=data,
        )


def _pdf(lines: int = 30, pages: int = 1) -> bytes:
    return images_to_pdf([text_page(lines=lines + i) for i in range(pages)])


@pytest.fixture
def env(tmp_path: Path) -> Iterator[Env]:
    settings = Settings(
        app_env="local",
        auth_mode="dev",
        bootstrap_admin_emails=ADMIN,
        data_dir=tmp_path,
        database_url=f"sqlite+aiosqlite:///{tmp_path / 'db.sqlite'}",
        render_dpi=72,
        llm_fixtures_dir=tmp_path / "fx",
        webhook_allow_private=True,  # test không phân giải DNS; kiểm tra SSRF có test riêng
    )
    engine = create_async_engine(settings.database_url)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    gp = sample_extraction()
    gp["so_gp"] = {"value": "635/GP-BTTTT", "confidence": "medium", "source_page": 1}
    provider = GuardProvider(
        MockProvider(responses={"GiayPhep": gp, "PhanLoai": {"loai_van_ban": "GP_HOAT_DONG_LUAT_2016"}})
    )

    async def session_override() -> AsyncIterator[AsyncSession]:
        async with maker() as s:
            yield s

    app = create_app()
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_session] = session_override

    async def enqueue_inline(doc_id: uuid.UUID, read_cache: bool = True, *, rerun: bool = False) -> None:
        async with maker() as s:
            await process_document(s, settings, doc_id, provider, read_cache=read_cache)

    keys: dict[str, str] = {}
    with TestClient(app) as client:

        async def init() -> None:
            async with engine.begin() as conn:
                await conn.run_sync(Base.metadata.create_all)
            async with maker() as s:
                s.add(models.Tenant(id=models.DEFAULT_TENANT_ID, slug="default", name="Nội bộ"))
                s.add(models.Tenant(id=TENANT_A, slug="a", name="Đối tác A", webhook_secret=SECRET_A))
                s.add(models.Tenant(id=TENANT_B, slug="b", name="Đối tác B", require_human_review=False))
                await s.flush()
                specs: dict[str, dict[str, Any]] = {
                    "a_live": {"tenant_id": TENANT_A, "sandbox": False},
                    "a_test": {"tenant_id": TENANT_A, "sandbox": True},
                    "b_live": {"tenant_id": TENANT_B, "sandbox": False},
                    "b_test": {"tenant_id": TENANT_B, "sandbox": True},
                    "a_read": {"tenant_id": TENANT_A, "sandbox": False, "scopes": ["documents:read"]},
                    "a_revoked": {"tenant_id": TENANT_A, "sandbox": False},
                    "a_expired": {"tenant_id": TENANT_A, "sandbox": False},
                }
                for name, spec in specs.items():
                    row, raw = new_api_key(settings, name=name, **spec)
                    if name == "a_revoked":
                        row.revoked_at = datetime.now(UTC)
                    if name == "a_expired":
                        row.expires_at = datetime.now(UTC) - timedelta(days=1)
                    s.add(row)
                    keys[name] = raw
                await s.commit()

        client.portal.call(init)  # type: ignore[union-attr]
        app.state.enqueue = enqueue_inline
        app.state.rate_limiter = MemoryRateLimiter()
        yield Env(
            client=client,
            settings=settings,
            provider=provider,
            maker=maker,
            keys=keys,
            hooks=[],
            hook_status=[],
        )
        client.portal.call(engine.dispose)  # type: ignore[union-attr]


def _count(model: Any, *where: Any):  # type: ignore[no-untyped-def]
    async def run(s: AsyncSession) -> int:
        return int(await s.scalar(select(func.count()).select_from(model).where(*where)) or 0)

    return run


# ---------------- Xác thực & phân quyền ----------------


def test_auth_errors(env: Env) -> None:
    c = env.client
    r = c.get("/v1/documents")
    assert r.status_code == 401 and r.json()["error"]["code"] == "missing_api_key"
    assert r.headers["WWW-Authenticate"] == "Bearer" and r.headers["X-Request-ID"].startswith("req_")
    for raw, code in [
        ("abc", "invalid_api_key"),
        ("ds_live_" + "x" * 32, "invalid_api_key"),
        (env.keys["a_revoked"], "api_key_revoked"),
        (env.keys["a_expired"], "api_key_expired"),
        (env.keys["a_live"].replace("ds_live_", "ds_test_"), "invalid_api_key"),  # đổi prefix -> sai hash
    ]:
        r = c.get("/v1/documents", headers={"Authorization": f"Bearer {raw}"})
        assert r.status_code == 401 and r.json()["error"]["code"] == code, raw
    r = env.upload("a_read", [("a.pdf", _pdf())])
    assert r.status_code == 403 and r.json()["error"]["code"] == "insufficient_scope"
    assert c.get("/v1/documents", headers=env.h("a_read")).status_code == 200


def test_missing_key_rejected_before_reading_body(env: Env) -> None:
    r = env.client.post("/v1/documents", files=[("files", ("a.pdf", _pdf(), "application/pdf"))])
    assert r.status_code == 401 and r.json()["error"]["code"] == "missing_api_key"
    assert env.db(_count(models.Document)) == 0


def test_validation_errors(env: Env) -> None:
    r = env.upload("a_live", [("a.pdf", b"not a pdf")])
    assert r.status_code == 415 and r.json()["error"]["code"] == "unsupported_media_type"
    r = env.upload("a_live", [("a.pdf", b"%PDF-1.4 broken")])
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_pdf"
    r = env.upload("a_live", [("a.pdf", _pdf(pages=31))])
    assert r.status_code == 422 and r.json()["error"]["code"] == "too_many_pages"
    r = env.upload("a_live", [("a.pdf", _pdf()), ("b.pdf", _pdf())], external_id="only-one")
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_request"
    r = env.upload("a_live", [("a.pdf", _pdf())], sandbox_scenario="failed")
    assert r.status_code == 422  # sandbox_scenario chỉ cho key ds_test_
    r = env.client.get("/v1/documents/not-a-uuid", headers=env.h("a_live"))
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_request"
    assert "details" in r.json()["error"]

    env.client.app.dependency_overrides[get_settings] = lambda: env.settings.model_copy(  # type: ignore[attr-defined]
        update={"max_file_mb": 0}
    )
    r = env.upload("a_live", [("a.pdf", _pdf())])
    assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"
    assert env.db(_count(models.Document)) == 0 and env.provider.calls == []


# ---------------- Cách ly tenant ----------------


def test_cross_tenant_access_is_blocked(env: Env) -> None:
    """Cố tình dùng key tenant B (và key sandbox của chính A) truy cập dữ liệu tenant A -> luôn 404."""
    r = env.upload("a_live", [("a.pdf", _pdf())], external_id="A-001", webhook_url="https://a.example/hook")
    assert r.status_code == 202, r.text
    doc_id, batch_id = r.json()["documents"][0]["id"], r.json()["batch_id"]
    assert env.client.get(f"/v1/documents/{doc_id}", headers=env.h("a_live")).status_code == 200

    for intruder in ("b_live", "b_test", "a_test"):
        h = env.h(intruder)
        for path in (
            f"/v1/documents/{doc_id}",
            f"/v1/documents/{doc_id}/export.xlsx",
            f"/v1/documents/{doc_id}/webhooks",
            f"/v1/batches/{batch_id}",
            f"/v1/batches/{batch_id}/export.zip",
        ):
            r = env.client.get(path, headers=h)
            assert r.status_code == 404 and r.json()["error"]["code"] == "not_found", (intruder, path)
        r = env.client.post(f"/v1/documents/{doc_id}/webhook/resend", headers=h)
        assert r.status_code == 404, intruder
        listing = env.client.get("/v1/documents", headers=h, params={"external_id": "A-001"}).json()
        assert listing["data"] == [], intruder
        # Không được chèn document vào batch của tenant khác
        r = env.upload(intruder, [("x.pdf", _pdf(lines=5))], batch_id=batch_id)
        assert r.status_code == 404, intruder
    usage_b = env.client.get("/v1/usage", headers=env.h("b_live")).json()
    assert usage_b["documents"] == 0 and usage_b["pages"] == 0


def test_same_file_other_tenant_is_processed_separately(env: Env) -> None:
    pdf = _pdf()
    a = env.upload("a_live", [("a.pdf", pdf)]).json()["documents"][0]
    calls = len(env.provider.calls)
    b = env.upload("b_live", [("a.pdf", pdf)]).json()["documents"][0]
    assert not b["deduplicated"] and len(env.provider.calls) > calls  # không dùng kết quả tenant khác
    assert a["id"] != b["id"]


# ---------------- Sandbox: 0 lần gọi LLM ----------------


def test_sandbox_never_calls_llm(env: Env) -> None:
    # Bật cả cờ gọi thật: sandbox vẫn không được chạm provider
    env.client.app.dependency_overrides[get_settings] = lambda: env.settings.model_copy(  # type: ignore[attr-defined]
        update={"allow_live_llm": True}
    )
    r = env.upload(
        "b_test", [("gp.pdf", _pdf(pages=2))], webhook_url="https://b.example/hook", external_id="T-1"
    )
    assert r.status_code == 202, r.text
    doc = r.json()["documents"][0]
    assert doc["sandbox"] is True
    got = env.client.get(f"/v1/documents/{doc['id']}", headers=env.h("b_test")).json()
    assert got["status"] == "completed" and got["needs_review"] is True  # B: không bắt buộc người duyệt
    assert got["result"]["so_gp"] == {
        "value": "635/GP-BTTTT",
        "confidence": "medium",
        "note": "số viết tay",
        "source_page": 1,
        "verified_by": None,
    }
    assert got["result"]["meta"]["file_name"] == "gp.pdf" and got["result"]["meta"]["pages"] == 2
    assert "provider" not in got["result"]["meta"]
    x = env.client.get(f"/v1/documents/{doc['id']}/export.xlsx", headers=env.h("b_test"))
    assert x.status_code == 200 and x.content[:2] == b"PK"

    # A bắt buộc người duyệt: sandbox mô phỏng đã duyệt
    a = env.upload("a_test", [("gp.pdf", _pdf(pages=2))]).json()["documents"][0]
    a = env.client.get(f"/v1/documents/{a['id']}", headers=env.h("a_test")).json()
    assert a["status"] == "completed" and a["needs_review"] is False

    for scenario in ("failed", "rejected"):
        d = env.upload(
            "a_test", [(f"{scenario}.pdf", _pdf(lines=3 + len(scenario)))], sandbox_scenario=scenario
        )
        d = env.client.get(f"/v1/documents/{d.json()['documents'][0]['id']}", headers=env.h("a_test")).json()
        assert d["status"] == scenario and d["result"] is None

    assert env.provider.calls == []
    assert env.db(_count(models.LlmCall)) == 0
    assert env.db(_count(models.WebhookDelivery)) == 1 and env.deliver() == 1
    usage = env.client.get("/v1/usage", headers=env.h("b_test")).json()
    assert (
        usage["mode"] == "sandbox"
        and usage["pages"] == 2
        and usage["llm_calls"] == 0
        and usage["cost_vnd"] == 0
    )
    # Dữ liệu sandbox không hiện ở giao diện rà soát nội bộ
    env.client.cookies.set(DEV_COOKIE, ADMIN)
    assert "gp.pdf" not in env.client.get("/documents").text


# ---------------- require_human_review ----------------


def test_require_human_review_true_releases_only_after_approval(env: Env) -> None:
    r = env.upload("a_live", [("gp.pdf", _pdf())], webhook_url="https://a.example/hook")
    doc_id = r.json()["documents"][0]["id"]
    assert len(env.provider.calls) > 0
    d = env.client.get(f"/v1/documents/{doc_id}", headers=env.h("a_live")).json()
    assert d["status"] == "pending_review" and d["result"] is None and d["needs_review"] is None
    x = env.client.get(f"/v1/documents/{doc_id}/export.xlsx", headers=env.h("a_live"))
    assert x.status_code == 409 and x.json()["error"]["code"] == "not_ready"
    assert env.db(_count(models.WebhookDelivery)) == 0  # chưa phải trạng thái cuối

    # Nhân viên nội bộ (tenant mặc định) thấy và duyệt document của đối tác trên web
    env.client.cookies.set(DEV_COOKIE, ADMIN)
    assert "gp.pdf" in env.client.get("/documents").text
    r = env.client.post(
        f"/documents/{doc_id}/save",
        data={"action": "approve", "f.so_gp": "636/GP-BTTTT"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    d = env.client.get(f"/v1/documents/{doc_id}", headers=env.h("a_live")).json()
    assert d["status"] == "completed" and d["needs_review"] is False
    assert d["result"]["so_gp"]["value"] == "636/GP-BTTTT"
    assert d["result"]["so_gp"]["verified_by"] == "ra_soat"  # không lộ email người duyệt
    assert ADMIN not in json.dumps(d)
    assert env.deliver() == 1
    body = json.loads(env.hooks[0].content)
    assert body["type"] == "document.completed" and body["data"]["document"]["id"] == doc_id


def test_require_human_review_false_returns_immediately(env: Env) -> None:
    r = env.upload("b_live", [("gp.pdf", _pdf())])
    d = env.client.get(f"/v1/documents/{r.json()['documents'][0]['id']}", headers=env.h("b_live")).json()
    assert d["status"] == "completed" and d["needs_review"] is True
    assert d["result"]["so_gp"]["confidence"] == "medium" and d["result_version"] == 1


# ---------------- Idempotency & chống trùng ----------------


def test_idempotency_key(env: Env) -> None:
    pdf = _pdf()
    r1 = env.upload("b_live", [("a.pdf", pdf)], idem="k-1", external_id="E1")
    r2 = env.upload("b_live", [("a.pdf", pdf)], idem="k-1", external_id="E1")
    assert r1.status_code == r2.status_code == 202
    assert r2.headers.get("Idempotent-Replayed") == "true" and r2.json() == r1.json()
    assert env.db(_count(models.Document)) == 1
    r3 = env.upload("b_live", [("a.pdf", _pdf(lines=7))], idem="k-1", external_id="E1")
    assert r3.status_code == 422 and r3.json()["error"]["code"] == "idempotency_key_reused"
    # Cùng key ở tenant khác là key khác
    r4 = env.upload("a_live", [("a.pdf", pdf)], idem="k-1", external_id="E1")
    assert r4.status_code == 202 and "Idempotent-Replayed" not in r4.headers
    # Request lỗi không bị lưu: sửa lỗi rồi gửi lại cùng key được
    bad = env.upload("b_live", [("a.pdf", b"nope")], idem="k-2")
    assert bad.status_code == 415
    ok = env.upload("b_live", [("a.pdf", _pdf(lines=9))], idem="k-2")
    assert ok.status_code == 202 and "Idempotent-Replayed" not in ok.headers


def test_same_file_same_tenant_not_reprocessed(env: Env) -> None:
    pdf = _pdf()
    first = env.upload("b_live", [("a.pdf", pdf)], external_id="X1").json()["documents"][0]
    calls = len(env.provider.calls)
    second = env.upload("b_live", [("a-copy.pdf", pdf)], external_id="X2", webhook_url="https://b.example/h")
    d2 = second.json()["documents"][0]
    assert len(env.provider.calls) == calls  # không gọi LLM lại
    assert d2["id"] != first["id"] and d2["deduplicated"] and d2["external_id"] == "X2"
    got1 = env.client.get(f"/v1/documents/{first['id']}", headers=env.h("b_live")).json()
    got2 = env.client.get(f"/v1/documents/{d2['id']}", headers=env.h("b_live")).json()
    assert got2["status"] == "completed" and got2["result"] == got1["result"]
    assert env.deliver() == 1  # bản trùng vẫn có webhook riêng
    usage = env.client.get("/v1/usage", headers=env.h("b_live")).json()
    assert usage["documents"] == 1  # bản trùng không tính hạn mức


def test_duplicate_follows_original_review(env: Env) -> None:
    pdf = _pdf()
    a1 = env.upload("a_live", [("a.pdf", pdf)]).json()["documents"][0]["id"]
    a2 = env.upload("a_live", [("a.pdf", pdf)], webhook_url="https://a.example/h").json()["documents"][0][
        "id"
    ]
    assert env.client.get(f"/v1/documents/{a2}", headers=env.h("a_live")).json()["status"] == "pending_review"
    env.client.cookies.set(DEV_COOKIE, ADMIN)
    env.client.post(f"/documents/{a1}/save", data={"action": "approve"}, follow_redirects=False)
    d2 = env.client.get(f"/v1/documents/{a2}", headers=env.h("a_live")).json()
    assert d2["status"] == "completed" and d2["result"] is not None
    assert env.deliver() == 1


# ---------------- Hạn mức ----------------


def _set_tenant(env: Env, tenant_id: uuid.UUID, **values: Any) -> None:
    async def run(s: AsyncSession) -> None:
        t = await s.get(models.Tenant, tenant_id)
        for k, v in values.items():
            setattr(t, k, v)
        await s.commit()

    env.db(run)


def test_rate_limit_per_minute(env: Env) -> None:
    _set_tenant(env, TENANT_B, rate_limit_per_minute=3)
    codes = [env.client.get("/v1/documents", headers=env.h("b_live")) for _ in range(4)]
    assert [r.status_code for r in codes] == [200, 200, 200, 429]
    assert codes[0].headers["X-RateLimit-Limit"] == "3" and codes[0].headers["X-RateLimit-Remaining"] == "2"
    assert codes[3].json()["error"]["code"] == "rate_limited" and int(codes[3].headers["Retry-After"]) > 0
    assert (
        env.client.get("/v1/documents", headers=env.h("a_live")).status_code == 200
    )  # key khác không ảnh hưởng


def test_monthly_page_quota_and_budget(env: Env) -> None:
    _set_tenant(env, TENANT_B, monthly_page_quota=3)
    assert env.upload("b_live", [("a.pdf", _pdf(pages=2))]).status_code == 202
    r = env.upload("b_live", [("b.pdf", _pdf(lines=11, pages=2))])
    assert r.status_code == 402 and r.json()["error"]["code"] == "page_quota_exceeded"
    assert env.upload("b_test", [("b.pdf", _pdf(lines=11, pages=2))]).status_code == 202  # sandbox miễn phí
    u = env.client.get("/v1/usage", headers=env.h("b_live")).json()
    assert u["pages"] == 2 and u["limits"]["pages_remaining"] == 1

    _set_tenant(env, TENANT_B, monthly_page_quota=None, monthly_budget_vnd=0.0)
    r = env.upload("b_live", [("c.pdf", _pdf(lines=13))])
    assert r.status_code in (202, 402)  # mock không có đơn giá -> ước tính 0 đ
    _set_tenant(env, TENANT_B, monthly_budget_vnd=-1.0)
    r = env.upload("b_live", [("d.pdf", _pdf(lines=14))])
    assert r.status_code == 402 and r.json()["error"]["code"] == "budget_exceeded"


# ---------------- Webhook ----------------


def test_webhook_signature_retry_and_resend(env: Env) -> None:
    r = env.upload("a_test", [("gp.pdf", _pdf())], webhook_url="https://a.example/hook", external_id="W-1")
    doc_id = r.json()["documents"][0]["id"]
    env.hook_status[:] = [500]
    assert env.deliver() == 0  # lần 1 lỗi -> chờ retry

    async def delivery(s: AsyncSession) -> models.WebhookDelivery:
        return (await s.scalars(select(models.WebhookDelivery))).one()

    d = env.db(delivery)
    assert d.status == "pending" and d.attempts == 1 and d.last_status_code == 500
    assert env.deliver() == 0  # chưa tới hạn (backoff 30s)

    async def make_due(s: AsyncSession) -> None:
        row = (await s.scalars(select(models.WebhookDelivery))).one()
        row.next_attempt_at = datetime.now(UTC) - timedelta(seconds=1)
        await s.commit()

    env.db(make_due)
    assert env.deliver() == 1
    req = env.hooks[-1]
    assert req.headers["X-DocSense-Event"] == "document.completed"
    assert webhooks.verify_signature(SECRET_A, req.content, req.headers["X-DocSense-Signature"])
    assert not webhooks.verify_signature("whsec_wrong", req.content, req.headers["X-DocSense-Signature"])
    assert not webhooks.verify_signature(SECRET_A, req.content + b" ", req.headers["X-DocSense-Signature"])
    payload = json.loads(req.content)
    assert payload["data"]["document"]["external_id"] == "W-1" and "result" not in payload["data"]["document"]

    hist = env.client.get(f"/v1/documents/{doc_id}/webhooks", headers=env.h("a_test")).json()
    assert [h["status"] for h in hist] == ["succeeded"] and hist[0]["attempts"] == 2

    r = env.client.post(
        f"/v1/documents/{doc_id}/webhook/resend",
        headers=env.h("a_test"),
        json={"webhook_url": "https://a.example/hook2"},
    )
    assert r.status_code == 202 and r.json()["status"] == "pending"
    assert env.deliver() == 1 and str(env.hooks[-1].url) == "https://a.example/hook2"
    assert json.loads(env.hooks[-1].content)["type"] == "document.completed"


def test_webhook_gives_up_after_max_attempts(env: Env) -> None:
    env.settings.webhook_max_attempts = 2
    env.upload("a_test", [("gp.pdf", _pdf())], webhook_url="https://a.example/hook")
    env.hook_status[:] = [503, 503]

    async def bump(s: AsyncSession) -> str:
        row = (await s.scalars(select(models.WebhookDelivery))).one()
        row.next_attempt_at = datetime.now(UTC) - timedelta(seconds=1)
        await s.commit()
        return row.status

    env.deliver()
    env.db(bump)
    env.deliver()
    assert env.db(bump) == "failed"


def test_webhook_not_final_resend_conflict(env: Env) -> None:
    r = env.upload("a_live", [("gp.pdf", _pdf())], webhook_url="https://a.example/hook")
    doc_id = r.json()["documents"][0]["id"]
    r = env.client.post(f"/v1/documents/{doc_id}/webhook/resend", headers=env.h("a_live"))
    assert r.status_code == 409 and r.json()["error"]["code"] == "not_final"


def test_webhook_url_ssrf_checks() -> None:
    s = Settings(webhook_allow_private=False)
    assert webhooks.validate_url("https://hooks.partner.vn/docsense", s) is None
    for bad in (
        "http://hooks.partner.vn/x",
        "https://localhost/x",
        "https://127.0.0.1/x",
        "https://10.0.0.5/x",
        "https://169.254.169.254/latest",
        "https://[::1]/x",
        "https://user:pw@hooks.partner.vn/x",
        "ftp://hooks.partner.vn/x",
    ):
        assert webhooks.validate_url(bad, s) is not None, bad


async def test_webhook_send_blocks_private_resolution() -> None:
    s = Settings(webhook_allow_private=False)
    d = models.WebhookDelivery(id=uuid.uuid4(), url="https://127.0.0.1/x", payload={}, event="e")
    sent: list[httpx.Request] = []
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(lambda r: sent.append(r) or httpx.Response(200))
    ) as c:
        code, err = await webhooks._send(d, "whsec_x", s, c)
    assert code is None and err is not None and sent == []


# ---------------- Batch ZIP ----------------


def test_zip_batch_flow(env: Env) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("BÁO A/gp1.pdf", _pdf(lines=21))
        zf.writestr("BÁO A/gp2.pdf", _pdf(lines=22))
        zf.writestr("BÁO B/gp3.pdf", _pdf(lines=23))
    r = env.client.post(
        "/v1/batches",
        headers=env.h("b_test"),
        files={"file": ("lo.zip", buf.getvalue(), "application/zip")},
        data={"webhook_url": "https://b.example/hook"},
    )
    assert r.status_code == 202, r.text
    batch_id = r.json()["batch_id"]
    b = env.client.get(f"/v1/batches/{batch_id}", headers=env.h("b_test")).json()
    assert b["status"] == "completed" and b["counts"] == {"completed": 3} and b["name"] == "lo.zip"
    assert {d["folder_name"] for d in b["documents"]} == {"BÁO A", "BÁO B"}
    z = env.client.get(f"/v1/batches/{batch_id}/export.zip", headers=env.h("b_test"))
    assert z.status_code == 200
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        assert sorted(zf.namelist()) == ["BÁO A/gp1.xlsx", "BÁO A/gp2.xlsx", "BÁO B/gp3.xlsx", "TongHop.xlsx"]
        ws = load_workbook(io.BytesIO(zf.read("BÁO A/gp1.xlsx")))["ThongTin"]
        assert all("sandbox/fixture" not in str(c.value) for row in ws.iter_rows() for c in row)
    assert env.deliver() == 4
    events = [json.loads(h.content)["type"] for h in env.hooks]
    assert events.count("document.completed") == 3 and events[-1] == "batch.completed"
    r = env.client.post(
        "/v1/batches",
        headers=env.h("b_test"),
        files={"file": ("x.zip", b"PK\x03\x04junk", "application/zip")},
    )
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_zip"


def test_multi_file_upload_with_external_ids_and_list(env: Env) -> None:
    r = env.upload(
        "b_test", [("a.pdf", _pdf(lines=31)), ("b.pdf", _pdf(lines=32))], external_id=["E-a", "E-b"]
    )
    assert r.status_code == 202
    docs = r.json()["documents"]
    assert [d["external_id"] for d in docs] == ["E-a", "E-b"] and len({d["batch_id"] for d in docs}) == 1
    lst = env.client.get("/v1/documents", headers=env.h("b_test"), params={"external_id": "E-b"}).json()[
        "data"
    ]
    assert [d["file_name"] for d in lst] == ["b.pdf"]
    lst = env.client.get("/v1/documents", headers=env.h("b_test"), params={"status": "completed"}).json()[
        "data"
    ]
    assert len(lst) == 2
    # Thêm vào batch đã có
    r = env.upload("b_test", [("c.pdf", _pdf(lines=33))], batch_id=docs[0]["batch_id"])
    assert r.json()["batch_id"] == docs[0]["batch_id"]


# ---------------- Schema, OpenAPI ----------------


def test_schema_and_openapi(env: Env) -> None:
    s = env.client.get("/v1/schema", headers=env.h("a_read")).json()
    assert s["schema_version"] == "v1" and "so_gp" in s["json_schema"]["properties"]
    assert "pending_review" in s["statuses"]
    spec = env.client.get("/v1/openapi.json").json()
    assert all(p.startswith("/v1/") for p in spec["paths"])
    assert "/v1/documents" in spec["paths"] and "/v1/usage" in spec["paths"]
    assert "ApiKey" in spec["components"]["securitySchemes"]
    assert env.client.get("/v1/docs").status_code == 200


def test_partner_docs_list_every_error_code() -> None:
    from app.api.partner.errors import ERROR_CODES

    doc = (Path(__file__).resolve().parent.parent / "docs" / "PARTNER_API.md").read_text(encoding="utf-8")
    missing = [c for c, (status, _) in ERROR_CODES.items() if f"| {status} | `{c}` |" not in doc]
    assert missing == []
    coll = (
        Path(__file__).resolve().parent.parent
        / "docs"
        / "postman"
        / "VTCDocSense_Partner_API.postman_collection.json"
    )
    assert json.loads(coll.read_text(encoding="utf-8"))["info"]["name"]


# ---------------- Log & retention ----------------


def _text_pdf(marker: str) -> bytes:
    """PDF 1 trang có lớp chữ chứa `marker`."""
    stream = f"BT /F1 24 Tf 72 720 Td ({marker}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf_content_never_logged(env: Env) -> None:
    marker = "BIMAT-NOIDUNG-PDF-7731"
    records: list[logging.LogRecord] = []

    class Grab(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            records.append(record)

    h = Grab(level=logging.DEBUG)
    root = logging.getLogger()
    old = root.level
    root.addHandler(h)
    root.setLevel(logging.DEBUG)
    try:
        pdf = _text_pdf(marker)
        r = env.upload("b_live", [("bimat.pdf", pdf)], external_id="L-1")
        assert r.status_code == 202, r.text
        env.upload("b_live", [("bimat.pdf", b"%PDF-1.4 " + marker.encode())])  # PDF hỏng chứa marker
        env.client.get(f"/v1/documents/{r.json()['documents'][0]['id']}", headers=env.h("b_live"))
    finally:
        root.removeHandler(h)
        root.setLevel(old)
    assert any(getattr(x, "path", None) == "/v1/documents" for x in records)  # có log truy cập
    from app.core.logging import JsonFormatter

    fmt = JsonFormatter()
    dump = "\n".join(fmt.format(x) for x in records)
    assert marker not in dump and "%PDF" not in dump and env.keys["b_live"] not in dump


def test_retention_purges_files_keeps_results(env: Env) -> None:
    r = env.upload("b_live", [("old.pdf", _pdf())])
    doc_id = uuid.UUID(r.json()["documents"][0]["id"])
    cache = env.settings.cache_dir / "ab"
    cache.mkdir(parents=True)
    (cache / "abc.json").write_text("{}")

    async def age_and_purge(s: AsyncSession) -> tuple[retention.PurgeResult, str]:
        d = await s.get(models.Document, doc_id)
        assert d is not None
        d.created_at = datetime.now(UTC) - timedelta(days=env.settings.retention_days + 1)
        s.add(models.IdempotencyKey(tenant_id=TENANT_B, key="old", fingerprint="x", created_at=d.created_at))
        await s.commit()
        future = datetime.now(UTC) + timedelta(days=1)  # cache mới tạo cũng coi là cũ
        res = await retention.purge(s, env.settings.model_copy(update={"retention_days": 0}), now=future)
        return res, d.storage_path

    res, path = env.db(age_and_purge)
    assert res.documents == 1 and res.cache_files == 1 and res.idempotency_keys == 1
    assert not Path(path).exists()
    d = env.client.get(f"/v1/documents/{doc_id}", headers=env.h("b_live")).json()
    assert d["status"] == "completed" and d["result"] is not None
