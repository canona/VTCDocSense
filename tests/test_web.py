"""Giao diện web M3 (AUTH_MODE=dev, SQLite, provider mock). Không gọi API LLM thật."""

import io
import uuid
import zipfile
from collections.abc import AsyncIterator, Iterator
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.api.main import create_app
from app.api.routes.internal import get_session
from app.core.config import Settings, get_settings
from app.db import models
from app.db.session import Base
from app.providers.mock import MockProvider
from app.services.documents import process_document
from app.web.auth import DEV_COOKIE, verify_cf_jwt
from tests.helpers import images_to_pdf, sample_extraction, text_page

ADMIN = "caonntt@gmail.com"


class Env(SimpleNamespace):
    client: TestClient
    settings: Settings
    provider: MockProvider
    maker: async_sessionmaker[AsyncSession]

    def login(self, email: str) -> None:
        self.client.cookies.set(DEV_COOKIE, email)

    def db(self, coro_fn):  # type: ignore[no-untyped-def]
        async def run():  # type: ignore[no-untyped-def]
            async with self.maker() as s:
                return await coro_fn(s)

        return self.client.portal.call(run)  # type: ignore[union-attr]


@pytest.fixture
def env(tmp_path: Path) -> Iterator[Env]:
    settings = Settings(
        app_env="local",
        auth_mode="dev",
        bootstrap_admin_emails=ADMIN,
        data_dir=tmp_path,
        database_url=f"sqlite+aiosqlite:///{tmp_path / 'db.sqlite'}",
        render_dpi=72,
        llm_fixtures_dir=tmp_path / "fx",
    )
    engine = create_async_engine(settings.database_url)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    gp = sample_extraction()
    gp["so_gp"] = {"value": "635/GP-BTTTT", "confidence": "medium", "source_page": 1}
    provider = MockProvider(
        responses={"GiayPhep": gp, "PhanLoai": {"loai_van_ban": "GP_HOAT_DONG_LUAT_2016"}}
    )

    async def session_override() -> AsyncIterator[AsyncSession]:
        async with maker() as s:
            yield s

    app = create_app()
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[get_session] = session_override

    async def enqueue_inline(doc_id: uuid.UUID, read_cache: bool = True, *, rerun: bool = False) -> None:
        async with maker() as s:
            await process_document(s, settings, doc_id, provider, read_cache=read_cache)

    with TestClient(app) as client:

        async def init() -> None:
            async with engine.begin() as conn:
                await conn.run_sync(Base.metadata.create_all)
            async with maker() as s:
                s.add(models.Tenant(id=models.DEFAULT_TENANT_ID, slug="default", name="Nội bộ"))
                await s.commit()

        client.portal.call(init)  # type: ignore[union-attr]
        app.state.enqueue = enqueue_inline
        yield Env(client=client, settings=settings, provider=provider, maker=maker)
        client.portal.call(engine.dispose)  # type: ignore[union-attr]


def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _upload_and_start(e: Env) -> list[str]:
    pdf = images_to_pdf([text_page(), text_page(lines=10)])
    data = _zip({"BÁO A/gp1.pdf": pdf, "BÁO A/gp2.pdf": images_to_pdf([text_page(lines=20)])})
    r = e.client.post("/upload", files={"files": ("lo.zip", data, "application/zip")}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("/confirm")
    page = e.client.get(r.headers["location"])
    assert page.status_code == 200 and "Ước tính" in page.text and "Xác nhận chạy" in page.text
    assert e.provider.calls == []  # chưa chạy trước khi xác nhận
    batch_id = r.headers["location"].split("/")[2]
    r = e.client.post(f"/batches/{batch_id}/start", follow_redirects=False)
    assert r.status_code == 303

    async def ids(s: AsyncSession) -> list[str]:
        docs = await s.scalars(select(models.Document).order_by(models.Document.file_name))
        return [str(d.id) for d in docs]

    return e.db(ids)  # type: ignore[no-any-return]


def test_login_redirect_and_bootstrap_admin(env: Env) -> None:
    r = env.client.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    env.login(ADMIN)
    r = env.client.get("/")
    assert r.status_code == 200 and "Tổng quan" in r.text and "Quản trị" in r.text


def test_review_edit_history_approve_export(env: Env) -> None:
    env.login(ADMIN)
    ids = _upload_and_start(env)
    doc_id = ids[0]

    page = env.client.get(f"/documents/{doc_id}")
    assert page.status_code == 200
    assert 'name="f.so_gp"' in page.text and "c-medium" in page.text  # ô vàng cho medium
    assert "Văn bản cùng thư mục báo" in page.text and "gp2.pdf" in page.text
    img = env.client.get(f"/documents/{doc_id}/pages/1.png")
    assert img.status_code == 200 and img.content[:4] == b"\x89PNG"

    r = env.client.post(
        f"/documents/{doc_id}/save",
        data={"action": "save", "f.so_gp": "636/GP-BTTTT", "f.ngay_cap": "29/09/2021", "c.ngay_cap": "1"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    hist = env.client.get(f"/documents/{doc_id}/history").text
    assert "636/GP-BTTTT" in hist and "635/GP-BTTTT" in hist and ADMIN in hist and "v2" in hist

    async def check(s: AsyncSession) -> tuple[str | None, int, str]:
        d = await s.get(models.Document, uuid.UUID(doc_id))
        n = len((await s.scalars(select(models.FieldReview))).all())
        assert d is not None
        return d.so_gp, n, d.status

    so_gp, n_reviews, _ = env.db(check)
    assert so_gp == "636/GP-BTTTT" and n_reviews == 1  # ngày cấp đã high -> xác nhận không tạo lịch sử

    bad = env.client.post(f"/documents/{doc_id}/save", data={"action": "save", "f.ngay_cap": "31/02/2021"})
    assert bad.status_code == 422 and "Ngày cấp không hợp lệ" in bad.text

    env.client.post(f"/documents/{doc_id}/save", data={"action": "approve"}, follow_redirects=False)
    assert env.db(check)[2] == "approved"

    z = env.client.get("/export.zip?only_approved=1")
    assert z.status_code == 200
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        names = zf.namelist()
        assert names == ["BÁO A/gp1.xlsx", "TongHop.xlsx"]
        ws = load_workbook(io.BytesIO(zf.read("TongHop.xlsx")))["TongHop"]
        assert ws.cell(row=2, column=4).value == "636/GP-BTTTT"
        assert ws.cell(row=2, column=8).value == "Đã duyệt"
    assert len(env.client.get("/documents").text.split('name="ids"')) == 3  # 2 document


def test_viewer_cannot_edit_and_next_doc(env: Env) -> None:
    env.login(ADMIN)
    ids = _upload_and_start(env)
    env.login("viewer@x.vn")
    assert "Chế độ chỉ xem" in env.client.get(f"/documents/{ids[0]}").text
    assert env.client.post(f"/documents/{ids[0]}/save", data={"action": "approve"}).status_code == 403
    assert env.client.get("/upload").status_code == 403
    assert env.client.get("/admin").status_code == 403
    r = env.client.get(f"/review/next?after={ids[0]}", follow_redirects=False)
    assert r.headers["location"] in (
        f"/documents/{ids[1]}",
        "/documents?flash=Kh%C3%B4ng+c%C3%B2n+document+c%E1%BA%A7n+r%C3%A0+so%C3%A1t",
    )


def test_rerun_requires_confirmation(env: Env) -> None:
    env.login(ADMIN)
    ids = _upload_and_start(env)
    calls = len(env.provider.calls)
    modal = env.client.get(f"/documents/{ids[0]}/rerun").text
    assert "Ước tính" in modal and "Xác nhận chạy lại" in modal
    assert env.client.post(f"/documents/{ids[0]}/rerun").status_code == 400
    env.client.post(f"/documents/{ids[0]}/rerun", data={"confirm": "1"}, follow_redirects=False)
    assert len(env.provider.calls) > calls


def test_admin_api_key_lifecycle(env: Env) -> None:
    env.login(ADMIN)
    env.client.get("/")
    r = env.client.post(
        "/admin/keys",
        data={
            "tenant_id": str(models.DEFAULT_TENANT_ID),
            "name": "Đối tác A",
            "sandbox": "1",
            "expires_days": "30",
        },
        follow_redirects=False,
    )
    key = r.headers["location"].split("new_key=")[1]
    assert key.startswith("ds_test_")

    async def keys(s: AsyncSession) -> list[models.ApiKey]:
        return list(await s.scalars(select(models.ApiKey)))

    (k,) = env.db(keys)
    assert key not in (k.key_hash, k.prefix) and k.sandbox and k.expires_at is not None
    env.client.post(f"/admin/keys/{k.id}/revoke", follow_redirects=False)
    assert env.db(keys)[0].revoked_at is not None
    env.client.post("/admin/users", data={"email": "rv@x.vn", "role": "reviewer"}, follow_redirects=False)
    admin_page = env.client.get("/admin").text
    assert "rv@x.vn" in admin_page and "api_key.revoke" in admin_page


def test_cf_access_jwt(tmp_path: Path) -> None:
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    s = Settings(cf_access_team_domain="team.cloudflareaccess.com", cf_access_aud="aud123", data_dir=tmp_path)
    jwks = SimpleNamespace(get_signing_key_from_jwt=lambda _t: SimpleNamespace(key=key.public_key()))
    now = datetime.now(UTC)

    def token(**kw: object) -> str:
        claims = {
            "email": "A@B.vn",
            "aud": "aud123",
            "iss": "https://team.cloudflareaccess.com",
            "exp": now + timedelta(minutes=5),
        }
        return jwt.encode(claims | kw, key, algorithm="RS256")

    assert verify_cf_jwt(token(), s, jwks) == "a@b.vn"  # type: ignore[arg-type]
    for bad in (token(aud="khac"), token(iss="https://evil"), token(exp=now - timedelta(minutes=1))):
        with pytest.raises(HTTPException) as ei:
            verify_cf_jwt(bad, s, jwks)  # type: ignore[arg-type]
        assert ei.value.status_code == 401


def test_dev_auth_blocked_in_production(env: Env) -> None:
    env.client.app.dependency_overrides[get_settings] = lambda: env.settings.model_copy(  # type: ignore[attr-defined]
        update={"app_env": "production"}
    )
    env.login(ADMIN)
    assert env.client.get("/").status_code == 500
