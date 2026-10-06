from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.cli import main as cli_main
from app.core.config import Settings
from app.models.schema import GiayPhepCore, LoaiVanBan, Meta, PdfType
from app.pipeline.export import FILL, to_xlsx
from app.pipeline.pdf import PdfError, load_pdf
from app.pipeline.preprocess import LogicalPage, PreprocessResult
from app.pipeline.run import process_pdf, process_preprocessed
from app.pipeline.validate import validate
from app.providers.mock import MockProvider
from tests.helpers import GP_2016_P1, GP_2016_P2, images_to_pdf, sample_extraction, text_page

SETTINGS = Settings(provider_max_retries=1, render_dpi=72)


def _scan_pdf(pages: int = 2) -> bytes:
    return images_to_pdf([text_page() for _ in range(pages)])


async def test_scan_pdf_uses_model_for_classification_and_images() -> None:
    provider = MockProvider(
        responses={
            "PhanLoai": {"loai_van_ban": "GP_HOAT_DONG_LUAT_2016", "ten_loai_giay_phep": None, "ly_do": "x"},
            "GiayPhep": sample_extraction(),
        }
    )
    gp = await process_pdf(_scan_pdf(), "scan.pdf", provider, SETTINGS)
    assert [c.schema_name for c in provider.calls] == ["PhanLoai", "GiayPhep"]
    extract_req = provider.calls[1]
    assert all(p.image and p.image_mime == "image/jpeg" and p.text is None for p in extract_req.pages)
    assert gp.meta.pdf_type == PdfType.SCAN
    assert gp.meta.classified_by == "model"
    assert gp.meta.provider == "mock"
    assert gp.loai_van_ban == LoaiVanBan.GP_HOAT_DONG_LUAT_2016
    assert gp.so_gp.value == "635/GP-BTTTT"


async def test_text_pages_rule_classified_and_page1_image_attached() -> None:
    img = text_page()
    pre = PreprocessResult(
        pages=[
            LogicalPage(page_no=1, source="1", text=GP_2016_P1, image=img),
            LogicalPage(page_no=2, source="2", text=GP_2016_P2 * 2),
        ],
        dropped=[],
    )
    provider = MockProvider(responses={"GiayPhep": sample_extraction(loai_van_ban="KHAC")})
    gp = await process_preprocessed(
        pre, file_name="t.pdf", physical_pages=2, kind=PdfType.TEXT, provider=provider, settings=SETTINGS
    )
    assert [c.schema_name for c in provider.calls] == ["GiayPhep"]  # không gọi model phân loại
    p1, p2 = provider.calls[0].pages
    assert p1.text and p1.image  # trang 1 kèm ảnh để đọc số/ngày trong chữ ký số
    assert p2.text and p2.image is None
    assert gp.meta.classified_by == "rule"
    assert gp.loai_van_ban == LoaiVanBan.GP_HOAT_DONG_LUAT_2016  # rule thắng output model


async def test_khac_skips_extraction() -> None:
    provider = MockProvider()  # mặc định phân loại KHAC
    gp = await process_pdf(_scan_pdf(1), "khac.pdf", provider, SETTINGS)
    assert [c.schema_name for c in provider.calls] == ["PhanLoai"]
    assert gp.loai_van_ban == LoaiVanBan.KHAC
    assert gp.needs_review


async def test_invalid_json_retries_then_fallback() -> None:
    bad = MockProvider(response={"loai_van_ban": "GP_HOAT_DONG_LUAT_2016", "an_pham": "không phải list"})
    good = MockProvider(responses={"GiayPhep": sample_extraction()})
    pre = PreprocessResult(pages=[LogicalPage(page_no=1, source="1", text=GP_2016_P1)], dropped=[])
    gp = await process_preprocessed(
        pre,
        file_name="t.pdf",
        physical_pages=1,
        kind=PdfType.TEXT,
        provider=bad,
        settings=SETTINGS,
        fallback=good,
    )
    assert len(bad.calls) == 2  # 1 lần + 1 retry
    assert len(good.calls) == 1
    assert gp.meta.provider == "mock" and gp.so_gp.value == "635/GP-BTTTT"


def test_pdf_errors() -> None:
    with pytest.raises(PdfError) as e:
        load_pdf(b"not a pdf", max_pages=30, max_mb=20)
    assert e.value.code == "corrupt"
    with pytest.raises(PdfError) as e:
        load_pdf(_scan_pdf(3), max_pages=2, max_mb=20)
    assert e.value.code == "too_many_pages"
    with pytest.raises(PdfError) as e:
        load_pdf(b"x" * (1024 * 1024 + 1), max_pages=30, max_mb=1)
    assert e.value.code == "too_large"


def test_excel_export_sheets_and_colors() -> None:
    gp = validate(
        GiayPhepCore.model_validate(sample_extraction()),
        Meta(file_name="a.pdf", pages=3, pdf_type=PdfType.TEXT),
    )
    wb = load_workbook(BytesIO(to_xlsx(gp, folder="BÁO AN GIANG")))
    assert wb.sheetnames == ["ThongTin", "AnPham", "LanhDao"]
    ws = wb["ThongTin"]
    rows = {r[0].value: r for r in ws.iter_rows() if r[0].value}
    assert rows["Số GP"][1].value == "635/GP-BTTTT"
    assert rows["Số GP"][2].value == "Trung bình"
    assert rows["Số GP"][1].fill.fgColor.rgb.endswith(FILL[gp.so_gp.confidence].fgColor.rgb[-6:])
    assert rows["Ngày cấp"][1].value == "29/09/2021"
    assert rows["Thư mục"][1].value == "BÁO AN GIANG"
    assert [c.value for c in wb["AnPham"][2]][:4] == [
        "635/GP-BTTTT",
        "Báo in",
        "Ấn phẩm chính",
        "Báo An Giang",
    ]
    assert [c.value for c in wb["LanhDao"][2]] == ["635/GP-BTTTT", "Tổng biên tập", "Trần Thị Bích Vân"]


def test_cli_extract_writes_outputs(tmp_path: Path) -> None:
    pdf = tmp_path / "gp.pdf"
    pdf.write_bytes(_scan_pdf(1))
    out = tmp_path / "out"
    assert cli_main(["extract", str(pdf), "--provider", "mock", "--out", str(out)]) == 0
    assert (out / "gp.json").exists() and (out / "gp.xlsx").exists()
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"broken")
    assert cli_main(["extract", str(bad), "--provider", "mock", "--out", str(out)]) == 1
