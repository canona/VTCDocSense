"""Xuất JSON + Excel (3 sheet ThongTin / AnPham / LanhDao), bố cục theo file đáp án mẫu."""

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.models.schema import ChucVu, Confidence, GiayPhep, LoaiHinh, LoaiVanBan, PdfType
from app.pipeline.validate import FIELD_LABELS, iter_fields

CONF_LABEL = {Confidence.high: "Cao", Confidence.medium: "Trung bình", Confidence.low: "Thấp"}
FILL = {
    Confidence.medium: PatternFill("solid", fgColor="FFF2CC"),  # vàng
    Confidence.low: PatternFill("solid", fgColor="F8CBAD"),  # cam
}
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
LOAI_LABEL = {
    LoaiVanBan.GP_HOAT_DONG_LUAT_1989: "Luật Báo chí 1989 (sửa đổi 1999)",
    LoaiVanBan.GP_HOAT_DONG_LUAT_2016: "Luật Báo chí 2016",
    LoaiVanBan.GP_MO_CHUYEN_TRANG: "Luật Báo chí 2016 (GP mở chuyên trang)",
    LoaiVanBan.KHAC: "Không nhận diện được",
}
PDF_TYPE_LABEL = {
    PdfType.TEXT: "PDF có sẵn chữ (lấy trực tiếp)",
    PdfType.SCAN: "Ảnh scan (AI đọc ảnh)",
    PdfType.MIXED: "Hỗn hợp chữ + ảnh scan",
}
LOAI_HINH_LABEL = {
    LoaiHinh.BAO_IN: "Báo in",
    LoaiHinh.BAO_DIEN_TU: "Báo điện tử",
    LoaiHinh.CHUYEN_TRANG: "Chuyên trang báo điện tử",
}
CHUC_VU_LABEL = {ChucVu.TONG_BIEN_TAP: "Tổng biên tập", ChucVu.PHO_TONG_BIEN_TAP: "Phó Tổng biên tập"}

AN_PHAM_COLS = [
    ("Loại hình", "loai_hinh"),
    ("Cấp ấn phẩm", "cap"),
    ("Tên gọi", "ten_goi"),
    ("Ngôn ngữ", "ngon_ngu"),
    ("Kỳ hạn xuất bản", "ky_han"),
    ("Thời gian phát hành", "thoi_gian_phat_hanh"),
    ("Khuôn khổ", "khuon_kho"),
    ("Số trang", "so_trang"),
    ("Số lượng", "so_luong"),
    ("Nơi in", "noi_in"),
    ("Tên miền", "ten_mien"),
    ("Đơn vị cung cấp Internet", "isp"),
]


def _fmt_date(d: date | None) -> str | None:
    return d.strftime("%d/%m/%Y") if d else None


def _header(ws: Worksheet, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL


def _thong_tin(ws: Worksheet, gp: GiayPhep, folder: str | None, extracted_on: date) -> None:
    ws.title = "ThongTin"
    ten_bao = gp.co_quan_bao_chi.ten.value or ""
    ws.append([f"TRÍCH XUẤT GIẤY PHÉP: {gp.so_gp.value or '(chưa rõ số)'} – {ten_bao}".rstrip(" –")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["File nguồn", gp.meta.file_name])
    ws.append(["Thư mục", folder or ""])
    ws.append(["Kiểu PDF", PDF_TYPE_LABEL[gp.meta.pdf_type]])
    ws.append(["Mẫu giấy phép", LOAI_LABEL[gp.loai_van_ban]])
    ws.append(["Ngày trích xuất", extracted_on.strftime("%d/%m/%Y")])
    _header(ws, ws.max_row + 1, ["Trường", "Giá trị", "Độ tin cậy", "Ghi chú kiểm tra"])

    rows: list[tuple[str, str | None, Confidence | None, str | None]] = [
        ("Loại văn bản", gp.ten_loai_giay_phep, Confidence.high if gp.ten_loai_giay_phep else None, None)
    ]
    for key, f in iter_fields(gp):
        value = _fmt_date(f.value) if key == "ngay_cap" else f.value
        note = f.note
        if f.verified_by:
            note = f"{note}; đối chiếu: {f.verified_by}" if note else f"đối chiếu: {f.verified_by}"
        rows.append((FIELD_LABELS[key], value, f.confidence if value or note else None, note))
    thay_the = "; ".join(f"{g.so_gp or '?'} ngày {g.ngay or '?'}" for g in gp.gp_duoc_thay_the)
    rows.append(("GP được thay thế", thay_the or None, Confidence.high if thay_the else None, None))

    for label, value, conf, note in rows:
        ws.append([label, value, CONF_LABEL[conf] if conf and value else None, note])
        if conf in FILL:
            for col in range(1, 5):
                ws.cell(row=ws.max_row, column=col).fill = FILL[conf]

    if gp.review_reasons:
        ws.append([])
        ws.append(["Cần rà soát", "\n".join(gp.review_reasons)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
    meta = gp.meta
    if meta.provider:  # bản cho đối tác (API /v1) đã bỏ thông tin provider/model
        ws.append([])
        ws.append(
            [
                "Xử lý",
                f"{meta.provider}/{meta.model} · {meta.duration_ms} ms · "
                f"{meta.input_tokens}+{meta.output_tokens} token · {meta.logical_pages}/{meta.pages} trang",
            ]
        )

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 50
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _an_pham(ws: Worksheet, gp: GiayPhep) -> None:
    _header(ws, 1, ["Số GP"] + [h for h, _ in AN_PHAM_COLS] + ["Ghi chú"])
    for ap in gp.an_pham:
        vals = [
            LOAI_HINH_LABEL[ap.loai_hinh] if k == "loai_hinh" else getattr(ap, k) for _, k in AN_PHAM_COLS
        ]
        ws.append([gp.so_gp.value, *vals, None])
    for col, w in zip(
        "ABCDEFGHIJKLMN", [16, 22, 15, 30, 14, 14, 24, 13, 10, 14, 16, 28, 30, 20], strict=True
    ):
        ws.column_dimensions[col].width = w


def _lanh_dao(ws: Worksheet, gp: GiayPhep) -> None:
    _header(ws, 1, ["Số GP", "Chức vụ", "Họ tên"])
    if not gp.lanh_dao:
        ws.append([gp.so_gp.value, "(GP không ghi lãnh đạo)", None])
    for ld in gp.lanh_dao:
        ws.append([gp.so_gp.value, CHUC_VU_LABEL[ld.chuc_vu], ld.ho_ten])
    for col, w in zip("ABC", [16, 22, 30], strict=True):
        ws.column_dimensions[col].width = w


def to_xlsx(gp: GiayPhep, *, folder: str | None = None, extracted_on: date | None = None) -> bytes:
    wb = Workbook()
    _thong_tin(wb.active, gp, folder, extracted_on or date.today())
    _an_pham(wb.create_sheet("AnPham"), gp)
    _lanh_dao(wb.create_sheet("LanhDao"), gp)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_json(gp: GiayPhep) -> str:
    return gp.model_dump_json(indent=2)


def write_outputs(gp: GiayPhep, out_dir: Path, stem: str, *, folder: str | None = None) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    json_path = out_dir / f"{stem}.json"
    xlsx_path = out_dir / f"{stem}.xlsx"
    json_path.write_text(to_json(gp), encoding="utf-8")
    xlsx_path.write_bytes(to_xlsx(gp, folder=folder))
    return [json_path, xlsx_path]
