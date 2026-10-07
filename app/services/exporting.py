"""Xuất ZIP giữ nguyên cấu trúc thư mục: <thư mục>/<tên PDF>.xlsx + TongHop.xlsx."""

import io
import zipfile
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.models import Document, Extraction
from app.models.schema import GiayPhep
from app.pipeline.export import LOAI_LABEL, to_xlsx
from app.services.partner import public_result

STATUS_LABEL = {
    "uploaded": "Đã tải lên",
    "processing": "Đang xử lý",
    "needs_review": "Cần rà soát",
    "auto_approved": "Tự động duyệt",
    "approved": "Đã duyệt",
    "rejected": "Từ chối",
    "failed": "Lỗi",
}
TONG_HOP_COLS = [
    ("Thư mục", 28),
    ("File", 40),
    ("Loại văn bản", 30),
    ("Số GP", 16),
    ("Ngày cấp", 12),
    ("Cơ quan chủ quản", 30),
    ("Cơ quan báo chí", 30),
    ("Trạng thái", 14),
    ("Người duyệt", 24),
]


def _stem(name: str) -> str:
    return name.rsplit(".", 1)[0]


def tong_hop(rows: list[tuple[Document, GiayPhep]], *, show_reviewer: bool = True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TongHop"
    ws.append([c for c, _ in TONG_HOP_COLS])
    for i, (_, w) in enumerate(TONG_HOP_COLS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        ws.column_dimensions[cell.column_letter].width = w
    for doc, gp in rows:
        ws.append(
            [
                doc.folder_name or "",
                doc.file_name,
                LOAI_LABEL[gp.loai_van_ban],
                gp.so_gp.value,
                gp.ngay_cap.value.strftime("%d/%m/%Y") if gp.ngay_cap.value else None,
                gp.co_quan_chu_quan.ten.value,
                gp.co_quan_bao_chi.ten.value,
                STATUS_LABEL.get(doc.status, doc.status),
                doc.reviewed_by if show_reviewer else None,
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_zip(items: list[tuple[Document, Extraction]], *, for_partner: bool = False) -> bytes:
    """for_partner: bỏ thông tin nội bộ (email người duyệt, provider/model) - dùng cho API /v1."""
    buf = io.BytesIO()
    rows: list[tuple[Document, GiayPhep]] = []
    used: set[str] = set()
    today = date.today()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for doc, ext in items:
            gp = GiayPhep.model_validate(public_result(ext.data) if for_partner else ext.data)
            rows.append((doc, gp))
            base = f"{doc.folder_name}/{_stem(doc.file_name)}" if doc.folder_name else _stem(doc.file_name)
            path, n = f"{base}.xlsx", 1
            while path in used:  # trùng tên (cùng file ở 2 batch)
                n += 1
                path = f"{base} ({n}).xlsx"
            used.add(path)
            zf.writestr(path, to_xlsx(gp, folder=doc.folder_name, extracted_on=today))
        zf.writestr("TongHop.xlsx", tong_hop(rows, show_reviewer=not for_partner))
    return buf.getvalue()
